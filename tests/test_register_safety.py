import io
import zipfile
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from iaa import cli
from iaa.core.config import Settings
from iaa.core.connectors import drive, eventbrite, mailchimp
from iaa.core.connectors.sheets import Workbook
from iaa.core.models import Proposal, Source, SourceRecord
from iaa.core.skill import InboxFile, Run
from iaa.core.tabular import read_rows
from iaa.registry import SKILLS
from iaa.skills.register import run as register_run
from iaa.skills.register.apply import apply
from iaa.skills.register.brief import render
from iaa.skills.register.models import KIND, Member
from iaa.skills.register.reconcile import reconcile, segment
from iaa.skills.register.run import _aliases, _read_sources, propose, refresh
from iaa.skills.register.sources import parse_insead_export

FIXTURES = Path(__file__).parent / "fixtures"
TODAY = date(2026, 11, 1)


def proposal(**changes):
    return Proposal(
        **(
            dict(
                run_id="test",
                kind=KIND,
                subject_id="1",
                subject="Synthetic Person",
                field="company",
                old="Old",
                new="New",
                reason="",
                source="",
                approve="yes",
            )
            | changes
        )
    )


def run_context(fake_book):
    settings = Settings(
        _env_file=None,
        register_sheet_id="synthetic",
        inbox_folder_id="synthetic",
        eventbrite_token="",
        mailchimp_api_key="",
        llm_provider="none",
    )
    return Run(Workbook(fake_book), settings)


def test_stale_and_repeated_approvals():
    member = Member(insead_id="1", name="Synthetic Person", company="Latest")
    updated, remaining, log = apply([member], [proposal()])
    assert updated == [member] and not log and "current value differs" in remaining[0].note
    updated, remaining, log = apply([member], [proposal(new="Latest")])
    assert updated == [member] and not remaining and log[0][-1] == "acknowledged"


@pytest.mark.parametrize(
    "field,new",
    [("compnay", "Example"), ("member_expire", "invalid"), ("status", "Actve"), ("events_registered", "-1")],
)
def test_invalid_manual_operations_stay_open(field, new):
    member = Member(insead_id="1", name="Synthetic Person")
    updated, remaining, log = apply([member], [proposal(field=field, new=new, old="")])
    assert updated == [member] and remaining and not log


def test_new_record_snapshot_cannot_change_after_review():
    member = Member(insead_id="1", name="Synthetic Person")
    pending = proposal(category="new-record", field="*", new=member.model_dump_json(), old="")
    updated, remaining, log = apply([], [pending], [member.model_copy(update={"company": "Changed"})])
    assert not updated and not log and "source record changed" in remaining[0].note


def test_every_exported_profile_field_is_reconciled():
    member = Member(insead_id="1", name="Synthetic Person")
    changes = dict(
        name="Synthetic Updated",
        naa="IAA SWE",
        programme="MBA",
        promotion="MBA'26",
        member_start="2026-01-01",
        email_insead="updated@example.org",
        industry="Consulting",
        function="Research",
        aliases=["alias@example.org"],
    )
    result = reconcile("test", [member], [member.model_copy(update=changes)], [], {}, as_of=TODAY)
    assert set(changes) <= {item.field for item in result.proposals}


def test_lapsed_segment_expires():
    member = Member(
        insead_id="1", name="Synthetic Person", status="Inactive", status_prev="Active", status_changed="2026-10-02"
    )
    assert segment(member, False, TODAY) == "lapsed-since-previous"
    assert segment(member, False, date(2027, 1, 1)) == "lapsed"


def test_cancelled_duplicate_and_recurring_events():
    member = Member(insead_id="1", name="Synthetic Person")
    record = SourceRecord(
        source=Source.EVENTBRITE,
        name=member.name,
        date=date(2026, 10, 1),
        payload={"event": "Synthetic event", "event_id": "1", "attendee_id": "1", "checked_in": False},
    )
    cancelled = record.model_copy(update={"payload": record.payload | {"attendee_id": "2", "cancelled": True}})
    recurring = record.model_copy(update={"date": date(2026, 10, 2)})
    result = reconcile("test", [member], None, [record, record, cancelled, recurring], {}, as_of=TODAY)
    assert len(result.aggregates["events"]) == 2
    assert all(counts["registered"] == 1 and counts["attended"] == 0 for counts in result.aggregates["events"].values())
    assert "| 0 |" in render(TODAY, result.aggregates)


def test_all_programmes_are_reported():
    members = [
        Member(insead_id=str(index), name="Synthetic Person", status="Active", programme=f"Programme {index}")
        for index in range(8)
    ]
    result = reconcile("test", members, None, [], {}, as_of=TODAY)
    assert sum(result.aggregates["by_programme_active"].values()) == 8


def test_renamed_event_has_one_aggregate():
    record = SourceRecord(
        source=Source.EVENTBRITE,
        name="Synthetic Person",
        date=date(2026, 10, 1),
        payload={"event_id": "1", "event": "Original title", "attendee_id": "1"},
    )
    renamed = record.model_copy(update={"payload": record.payload | {"event": "New title", "attendee_id": "2"}})
    result = reconcile("test", [], None, [record, renamed], {}, as_of=TODAY)
    assert len(result.aggregates["events"]) == 1
    assert next(iter(result.aggregates["events"].values()))["registered"] == 2


def test_existing_contact_suggestion_is_suppressed(fake_book):
    run = run_context(fake_book)
    run.book.write("members", Member, [Member(insead_id="1", name="Synthetic Person", do_not_contact=True)])
    assert refresh(run, [proposal(suggested_action="Renewal reminder")])[0].suggested_action.startswith("None:")


@pytest.mark.parametrize("decision", ["yes", "no"])
def test_existing_member_decisions_do_not_read_sources(fake_book, monkeypatch, decision):
    run = run_context(fake_book)
    run.book.write("members", Member, [Member(insead_id="1", name="Synthetic Person", company="Old")])
    run.book.write_proposals([proposal(approve=decision)])
    monkeypatch.setattr(drive, "list_inbox", lambda *args: pytest.fail("unexpected source read"))
    cli.apply_decisions(run, SKILLS)
    assert not run.book.proposals()


def test_newest_invalid_export_suspends_export_changes(fake_book):
    run = run_context(fake_book)
    raw = (FIXTURES / "insead_export.csv").read_bytes()
    run.inbox = [
        InboxFile("old.csv", raw, "old", "1", "seen"),
        InboxFile("new.csv", b"Download Alumni Details\ninvalid\n", "new", "2", "new"),
    ]
    box = _read_sources(run)
    assert box.export is None and box.export_error
    assert run.inbox[-1].kind == "insead-unreadable"


def test_other_files_in_the_inbox_do_not_block_the_export(fake_book):
    run = run_context(fake_book)
    document = io.BytesIO()
    with zipfile.ZipFile(document, "w") as archive:  # a .docx is a zip, like an .xlsx
        archive.writestr("word/document.xml", "<w:document/>")
    workbook = io.BytesIO()
    book = openpyxl.Workbook()
    book.active.append(["Budget", "SEK"])
    book.save(workbook)
    run.inbox = [
        InboxFile("export.csv", (FIXTURES / "insead_export.csv").read_bytes(), "a", "1", "new"),
        InboxFile("agenda.docx", document.getvalue(), "b", "2", "new"),
        InboxFile("budget.xlsx", workbook.getvalue(), "c", "3", "new"),
        InboxFile("notes.csv", b"a,b\n1,2\n", "d", "4", "new"),
    ]
    box = _read_sources(run)
    assert box.export and not box.export_error and box.unrecognised == 3
    assert [item.kind for item in run.inbox] == ["insead", "unknown", "unknown", "unknown"]
    proposals, brief = propose(run)
    assert proposals and "not recognised and not used: 3" in brief and "suspended" not in brief


def test_two_alias_rows_for_one_name_reach_the_matcher_as_a_conflict(fake_book):
    run = run_context(fake_book)
    members = [
        Member(insead_id="1", name="Synthetic One", status="Inactive", member_expire="2015-12-31"),
        Member(insead_id="2", name="Synthetic Two", status="Inactive", member_expire="2015-12-31"),
    ]
    run.book.append("aliases", ["alias", "insead_id"], [["Synthetic Payer", "1"], ["Synthetic Payer", 2]])
    assert _aliases(run) == {"synthetic payer": {"1", "2"}}
    payment = SourceRecord(
        source=Source.PAYMENTS, external_id="row1", name="Synthetic Payer", date=date(2026, 2, 1), status="Swish"
    )
    (row,) = [
        p
        for p in reconcile("test", members, None, [payment], _aliases(run), as_of=TODAY).proposals
        if p.field == "payment"
    ]
    assert (row.category, row.subject_id) == ("payment-unconfirmed-match", "")
    assert "aliases tab gives this name to several records" in row.reason and "['1', '2']" in row.reason


def test_a_newer_workbook_that_will_not_open_suspends_the_export(fake_book):
    run = run_context(fake_book)
    raw = (FIXTURES / "insead_export.csv").read_bytes()
    workbook = io.BytesIO()
    book = openpyxl.Workbook()
    for row in read_rows(raw):
        book.active.append(row)
    book.save(workbook)
    cut_short = workbook.getvalue()[: len(workbook.getvalue()) // 2]
    run.inbox = [InboxFile("old.csv", raw, "a", "1", "seen"), InboxFile("new.xlsx", cut_short, "b", "2", "new")]
    box = _read_sources(run)
    assert box.export is None and box.export_error and not box.unrecognised
    assert run.inbox[-1].kind == "unreadable"
    assert "export changes are suspended" in propose(run)[1]
    # the other way round, the valid export is the newer one and is used
    run = run_context(fake_book)
    run.inbox = [InboxFile("old.xlsx", cut_short, "b", "1", "seen"), InboxFile("new.csv", raw, "a", "2", "new")]
    box = _read_sources(run)
    assert box.export and not box.export_error


def test_the_inbox_is_parsed_once_however_many_new_records_are_approved(fake_book, monkeypatch):
    run = run_context(fake_book)
    raw = (FIXTURES / "insead_export.csv").read_bytes()
    run.inbox = [
        InboxFile("old.csv", raw, "a", "1", "seen"),
        InboxFile("new.csv", b"Download Alumni Details\ninvalid\n", "b", "2", "new"),
    ]
    calls = []
    monkeypatch.setattr(
        register_run, "_read_sources", lambda *args, **kwargs: calls.append(1) or _read_sources(*args, **kwargs)
    )
    applier = SKILLS[0].applier(run)
    approved = [
        item.model_copy(update={"approve": "yes"})
        for item in reconcile("test", [], parse_insead_export(raw), [], {}, as_of=TODAY).proposals
    ]
    assert len(approved) == 30
    assert {applier.carry_out(item) for item in approved} == {("", "the latest export is unavailable or invalid")}
    assert len(calls) == 1


def test_download_failure_is_recorded_and_blocks_stale_export(fake_book, monkeypatch):
    run = run_context(fake_book)
    monkeypatch.setattr(
        drive,
        "list_inbox",
        lambda *args: [
            {"id": "1", "name": "synthetic.csv", "mimeType": "text/csv"},
            {"id": "2", "name": "folder", "mimeType": "application/vnd.google-apps.folder"},
        ],
    )
    monkeypatch.setattr(drive, "download", lambda *args: (_ for _ in ()).throw(RuntimeError("synthetic")))
    assert [item.state for item in run.inbox] == ["unreadable", "unsupported"]
    assert _read_sources(run).export_error


def test_optional_connectors_are_used_only_for_reconciliation(fake_book, monkeypatch):
    run = run_context(fake_book)
    run.inbox = []
    run.settings.eventbrite_token = run.settings.eventbrite_org_id = "synthetic"
    run.settings.mailchimp_api_key = run.settings.mailchimp_list_id = "synthetic"
    calls = []
    monkeypatch.setattr(eventbrite, "fetch_attendees", lambda *args: calls.append("eventbrite") or [])
    monkeypatch.setattr(mailchimp, "fetch_audience", lambda *args: calls.append("mailchimp") or [])
    _read_sources(run, include_records=False)
    assert not calls
    _read_sources(run)
    assert calls == ["eventbrite", "mailchimp"]


def test_activity_fields_are_proposed_and_applied():
    members = parse_insead_export((FIXTURES / "insead_export.csv").read_bytes())
    record = SourceRecord(
        source=Source.EVENTBRITE,
        name=members[0].name,
        emails=["new@example.org"],
        date=date(2026, 10, 1),
        payload={"event_id": "1", "checked_in": True},
    )
    result = reconcile("test", members, None, [record], {}, as_of=TODAY)
    approved = [item.model_copy(update={"approve": "yes"}) for item in result.proposals]
    updated, remaining, _ = apply(members, approved)
    assert not remaining
    member = updated[0]
    assert member.events_registered == member.events_attended == 1
    assert member.segment == "paying-active" and "new@example.org" in member.aliases
